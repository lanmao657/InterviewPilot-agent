import base64
import logging
from io import BytesIO

from docx import Document as DocxDocument
from fastapi import UploadFile
from openai import AsyncOpenAI
from pptx import Presentation
from pypdf import PdfReader
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.models import Document, DocumentChunk
from app.services.embedding import EmbeddingService

logger = logging.getLogger(__name__)

# pypdf 提取结果低于此字符数时，判定为扫描件，走 OCR 回退
_OCR_THRESHOLD = 50


async def _ocr_pdf_with_vision(data: bytes) -> str:
    """用 pymupdf 渲染 PDF 页面为图片，再调 OpenAI Vision API 提取文字"""
    import pymupdf  # type: ignore[import-untyped]

    settings = get_settings()
    client = AsyncOpenAI(
        base_url=settings.ai_base_url,
        api_key=settings.ai_api_key or "",
    )

    doc = pymupdf.open(stream=data, filetype="pdf")
    page_texts: list[str] = []

    for page in doc:
        # 渲染为 PNG（300 DPI 保证清晰度）
        pix = page.get_pixmap(dpi=300)
        img_bytes = pix.tobytes("png")
        b64 = base64.b64encode(img_bytes).decode("utf-8")

        resp = await client.chat.completions.create(
            model=settings.ai_fast_model or settings.ai_model,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {
                            "type": "text",
                            "text": (
                                "请提取这张图片中的所有文字内容，保持原始排版结构，"
                                "直接输出文字，不要添加任何解释。"
                            ),
                        },
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:image/png;base64,{b64}",
                            },
                        },
                    ],
                }
            ],
            max_tokens=4096,
        )
        page_texts.append(resp.choices[0].message.content or "")

    doc.close()
    return "\n".join(page_texts).strip()


async def extract_upload_text(file: UploadFile) -> str:
    data = await file.read()
    filename = file.filename or "uploaded-file"
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if suffix == "pdf":
        reader = PdfReader(BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()

        # pypdf 提取不到足够文字时，用 Vision API 做 OCR 回退
        if len(text) < _OCR_THRESHOLD:
            logger.info("pypdf 提取结果不足 %d 字符，尝试 OCR 回退", _OCR_THRESHOLD)
            try:
                text = await _ocr_pdf_with_vision(data)
            except Exception:
                logger.exception("OCR 回退失败")
                # 如果 OCR 也失败，返回 pypdf 的原始结果（可能为空）
        return text
    if suffix == "docx":
        doc = DocxDocument(BytesIO(data))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs).strip()
    if suffix == "pptx":
        prs = Presentation(BytesIO(data))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        return "\n".join(texts).strip()
    return data.decode("utf-8", errors="ignore").strip()


def summarize_document(text: str, kind: str) -> dict:
    """生成文档摘要（保留向后兼容）"""
    compact = " ".join(text.split())[:600]
    if kind == "resume":
        return {
            "type": "简历",
            "highlights": ["项目经历", "核心技能", "过往职责"],
            "preview": compact or "未解析到文本，请检查文件内容。",
        }
    return {
        "type": "JD",
        "highlights": ["岗位职责", "能力要求", "面试关注点"],
        "preview": compact or "未解析到文本，请检查文件内容。",
    }


class DocumentService:
    def __init__(self, embedding_service: EmbeddingService, db: Session):
        self.embedding_service = embedding_service
        self.db = db

    async def process_document(self, document_id: int) -> None:
        """处理文档：切片 + 向量化"""
        # 获取文档
        result = self.db.execute(
            select(Document).where(Document.id == document_id)
        )
        document = result.scalar_one_or_none()
        if not document:
            return

        # 先删除已有的切片（幂等性保护）
        self.db.execute(
            delete(DocumentChunk).where(DocumentChunk.document_id == document_id)
        )

        # 切片
        chunks = self.embedding_service.chunk_text(document.content)

        # 生成 embeddings
        embeddings = await self.embedding_service.embed(chunks)

        # 保存切片到数据库
        for i, (chunk_content, embedding) in enumerate(zip(chunks, embeddings)):
            chunk = DocumentChunk(
                document_id=document_id,
                user_id=document.user_id,
                chunk_index=i,
                content=chunk_content,
                embedding=embedding,
            )
            self.db.add(chunk)

        self.db.commit()

    def summarize_document(self, text: str, kind: str) -> dict:
        """生成文档摘要"""
        return summarize_document(text, kind)
